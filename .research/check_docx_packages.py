from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

for name in ['paper/ARIA_complete_research_paper.docx','.research/topic_dossier.docx']:
    p=Path(name); doc=Document(p)
    with ZipFile(p) as z:
        names=set(z.namelist()); xml=z.read('word/document.xml').decode('utf-8','replace')
        media=[x for x in names if x.startswith('word/media/')]
        comments=[x for x in names if 'comment' in x.lower()]
        tracked=xml.count('<w:ins')+xml.count('<w:del')
        fields=xml.count('<w:fldSimple')+xml.count('<w:instrText')
    texts=[x.text.strip() for x in doc.paragraphs if x.text.strip()]
    tables=sum(len(t.rows) for t in doc.tables)
    missing=[]
    for phrase in (['33-dimensional','0.9121','0.4618','research prototype'] if name.startswith('paper') else ['Worth pursuing','249','Zotero']):
        if phrase not in '\n'.join(texts) and not any(phrase in c.text for t in doc.tables for r in t.rows for c in r.cells): missing.append(phrase)
    print(f'{name}: paragraphs={len(doc.paragraphs)} nonempty={len(texts)} tables={len(doc.tables)} table_rows={tables} media={len(media)} comments={len(comments)} tracked_changes={tracked} fields={fields} missing_required={missing} size={p.stat().st_size}')
