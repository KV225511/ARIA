from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from PIL import Image as PILImage

ROOT=Path(__file__).resolve().parents[1]

def clean(s):
    s=re.sub(r'!\[([^]]*)\]\([^)]*\)',r'\1',s)
    s=re.sub(r'\*\*([^*]+)\*\*',r'\1',s); s=re.sub(r'`([^`]+)`',r'\1',s)
    s=s.replace('  ',' ')
    return s.strip()

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def add_page_number(section):
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)

def add_docx_paragraph(doc,text,style=None):
    p=doc.add_paragraph(style=style); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    parts=re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'): p.add_run(part[2:-2]).bold=True
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]); r.font.name='Consolas'; r.font.size=Pt(9)
        else: p.add_run(part)
    return p

def docx_from_markdown(src,dst,title_page=False):
    lines=src.read_text(encoding='utf-8').splitlines(); doc=Document()
    sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    styles=doc.styles
    styles['Normal'].font.name='Times New Roman'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(4); styles['Normal'].paragraph_format.line_spacing=1.08
    for n,size,col in [('Title',20,'17365D'),('Heading 1',15,'17365D'),('Heading 2',12,'2F5D7C'),('Heading 3',11,'2F5D7C')]:
        st=styles[n]; st.font.name='Arial'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(col); st.font.bold=True
    add_page_number(sec)
    i=0; first_title=True
    while i<len(lines):
        raw=lines[i].rstrip(); s=raw.strip()
        if not s: i+=1; continue
        if s.startswith('|') and i+1<len(lines) and re.match(r'^\|?\s*:?-+',lines[i+1].strip()):
            rows=[]; i+=2
            header=[clean(x) for x in s.strip('|').split('|')]
            rows.append(header)
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([clean(x) for x in lines[i].strip().strip('|').split('|')]); i+=1
            cols=max(len(r) for r in rows); table=doc.add_table(rows=len(rows),cols=cols); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
            for ri,row in enumerate(rows):
                for ci,val in enumerate(row):
                    cell=table.cell(ri,ci); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text=val
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                        for r in p.runs: r.font.name='Arial'; r.font.size=Pt(8); r.bold=(ri==0)
                    if ri==0: set_cell_shading(cell,'DCE6F1')
            continue
        m=re.match(r'^(#{1,3})\s+(.*)',s)
        if m:
            level=len(m.group(1)); txt=clean(m.group(2))
            if level==1 and first_title:
                p=doc.add_paragraph(txt,style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; first_title=False
            else:
                if txt=='References': doc.add_page_break()
                doc.add_paragraph(txt,style=f'Heading {min(level,3)}')
            i+=1; continue
        im=re.match(r'^!\[(.*?)\]\((.*?)\)$',s)
        if im:
            path=(src.parent/im.group(2)).resolve(); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(6.6)); cap=doc.add_paragraph(im.group(1)); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.runs[0].italic=True; cap.runs[0].font.size=Pt(9); i+=1; continue
        if s.startswith('**') and s.endswith('**  '): s=s[:-2]
        if re.match(r'^[-*]\s+',s): add_docx_paragraph(doc,re.sub(r'^[-*]\s+','',s),'List Bullet'); i+=1; continue
        if re.match(r'^\d+\.\s+',s): add_docx_paragraph(doc,re.sub(r'^\d+\.\s+','',s),'List Number'); i+=1; continue
        p=add_docx_paragraph(doc,s)
        if first_title is False and len(doc.paragraphs)<=5: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        i+=1
    props=doc.core_properties; props.title=clean(lines[0].lstrip('# ')); props.author='Raghav Sejpal; Krissh Verma'; props.subject='ARIA research manuscript'; props.keywords='ARIA, adaptive interviewing, offline reinforcement learning, auditability'
    dst.parent.mkdir(parents=True,exist_ok=True); doc.save(dst)

styles=getSampleStyleSheet()
PDF_BODY=ParagraphStyle('Body',parent=styles['BodyText'],fontName='Times-Roman',fontSize=9.0,leading=11.5,alignment=TA_JUSTIFY,spaceAfter=4.5)
PDF_H1=ParagraphStyle('H1',parent=styles['Heading1'],fontName='Helvetica-Bold',fontSize=15,leading=18,textColor=colors.HexColor('#17365D'),spaceBefore=10,spaceAfter=6)
PDF_H2=ParagraphStyle('H2',parent=styles['Heading2'],fontName='Helvetica-Bold',fontSize=11.5,leading=14,textColor=colors.HexColor('#2F5D7C'),spaceBefore=8,spaceAfter=4)
PDF_TITLE=ParagraphStyle('Title',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=19,leading=23,textColor=colors.HexColor('#17365D'),alignment=TA_CENTER,spaceAfter=14)
PDF_SMALL=ParagraphStyle('Small',parent=PDF_BODY,fontSize=7.1,leading=8.6)

def esc(s):
    import html
    s=html.escape(s); s=re.sub(r'\*\*([^*]+)\*\*',r'<b>\1</b>',s); s=re.sub(r'`([^`]+)`',r'<font name="Courier">\1</font>',s); return s

def footer(canvas,doc):
    canvas.saveState(); canvas.setFont('Helvetica',8); canvas.setFillColor(colors.HexColor('#5A6872')); canvas.drawCentredString(A4[0]/2,0.42*inch,str(doc.page)); canvas.restoreState()

def pdf_from_markdown(src,dst):
    lines=src.read_text(encoding='utf-8').splitlines(); story=[]; i=0; first=True
    while i<len(lines):
        s=lines[i].strip()
        if not s: story.append(Spacer(1,3)); i+=1; continue
        if s.startswith('|') and i+1<len(lines) and re.match(r'^\|?\s*:?-+',lines[i+1].strip()):
            data=[[clean(x) for x in s.strip('|').split('|')]]; i+=2
            while i<len(lines) and lines[i].strip().startswith('|'):
                data.append([clean(x) for x in lines[i].strip().strip('|').split('|')]); i+=1
            n=max(len(r) for r in data); data=[r+['']*(n-len(r)) for r in data]
            pdata=[[Paragraph(esc(x),PDF_SMALL) for x in r] for r in data]
            widths=[(A4[0]-1.3*inch)/n]*n
            t=Table(pdata,colWidths=widths,repeatRows=1,hAlign='CENTER'); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE6F1')),('TEXTCOLOR',(0,0),(-1,0),colors.HexColor('#17365D')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#8795A1')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)])); story.extend([t,Spacer(1,5)]); continue
        m=re.match(r'^(#{1,3})\s+(.*)',s)
        if m:
            lvl=len(m.group(1)); txt=esc(clean(m.group(2)))
            if clean(m.group(2))=='References': story.append(PageBreak())
            story.append(Paragraph(txt,PDF_TITLE if lvl==1 and first else (PDF_H1 if lvl<=2 else PDF_H2))); first=False; i+=1; continue
        im=re.match(r'^!\[(.*?)\]\((.*?)\)$',s)
        if im:
            path=(src.parent/im.group(2)).resolve()
            with PILImage.open(path) as pimg: iw,ih=pimg.size
            scale=min((6.9*inch)/iw,(4.7*inch)/ih)
            pic=Image(str(path),width=iw*scale,height=ih*scale); pic.hAlign='CENTER'; story.extend([pic,Paragraph(esc(im.group(1)),ParagraphStyle('Caption',parent=PDF_SMALL,alignment=TA_CENTER,fontName='Times-Italic')),Spacer(1,6)]); i+=1; continue
        if re.match(r'^[-*]\s+',s): story.append(Paragraph('• '+esc(re.sub(r'^[-*]\s+','',s)),PDF_BODY)); i+=1; continue
        if re.match(r'^\d+\.\s+',s): story.append(Paragraph(esc(s),PDF_BODY)); i+=1; continue
        story.append(Paragraph(esc(s),PDF_BODY)); i+=1
    dst.parent.mkdir(parents=True,exist_ok=True)
    doc=SimpleDocTemplate(str(dst),pagesize=A4,rightMargin=.65*inch,leftMargin=.65*inch,topMargin=.62*inch,bottomMargin=.62*inch,title=clean(lines[0].lstrip('# ')),author='Raghav Sejpal; Krissh Verma')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

docx_from_markdown(ROOT/'.research/topic_dossier.md',ROOT/'.research/topic_dossier.docx')
docx_from_markdown(ROOT/'paper/ARIA_complete_research_paper.md',ROOT/'paper/ARIA_complete_research_paper.docx')
pdf_from_markdown(ROOT/'paper/ARIA_complete_research_paper.md',ROOT/'paper/ARIA_complete_research_paper.pdf')
pdf_from_markdown(ROOT/'.research/topic_dossier.md',ROOT/'.tmp/topic_dossier_layout_preview.pdf')
